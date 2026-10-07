"""Additive dashboard data; copy into the mentorship app.

Uses the supplied mentorship models. English topic text is the default topic
description, with Dari/Pashto fallback. No schema changes are required.
"""
from collections import Counter, defaultdict
from itertools import groupby
from statistics import mean, median

from django.conf import settings
from django.db.models import Count, F, Q

from .models import MenteeTopicStatus, MentorshipTopics, Staff
from .dashboard_quality import build_quality_data


def mentor_visit_counts(details_qs):
    """Match the original 'visit by mentors' export: date + normalized name.

    Repeated details/visit IDs for the same mentor on the same date count once.
    Missing mentor names/dates do not contribute. Facility is not in the key.
    """
    sets = defaultdict(lambda: defaultdict(set))
    global_keys = set()
    rows = details_qs.order_by().values_list(
        'mentorshipvistfk__visitdate', 'mentor__name',
        'mentorshipvistfk__facilityfk__districtfk__provincefk__name',
        'mentorshipvistfk__facilityfk__districtfk__name',
        'mentorshipvistfk__facilityfk__name', 'mentorshipvistfk__facilityfk__hfcode',
        'mentorshipvistfk__facilityfk_id', 'thematicname_id', 'topicname_id',
        'menteename__gender', 'menteename_id',
    )
    for day, mentor, province, district, facility, code, fid, thematic, topic, gender, mentee in rows.iterator(chunk_size=2000):
        name = (mentor or '').strip().lower()
        if not day or not name:
            continue
        key = (day, name)
        global_keys.add(key)
        month = (day.year, day.month)
        for group, bucket in (
            ('province', province), ('facility', (province, district, facility, code)),
            ('month', month), ('province_month', (province, month)),
            ('thematic', thematic), ('topic', (thematic, topic)),
            ('gender', gender), ('mentee_facility', (fid, mentee)),
        ):
            if group in ('gender', 'mentee_facility') and mentee is None:
                continue
            sets[group][bucket].add(key)
    return {'total': len(global_keys), **{group: {bucket: len(keys) for bucket, keys in buckets.items()} for group, buckets in sets.items()}}


def summarize_learning_history(events, selected_pairs, filters, status_map=None):
    """Count distinct LS visits before first dated PC/MC, per mentee/topic.

    Same-day LS/competency chronology is unknown from dates alone. Such pairs,
    and pairs with undated records, are reported but excluded from statistics.
    Duplicate detail lines in one visit count as one learning session.
    """
    status_map = status_map or {}
    groups = defaultdict(list)
    for event in events:
        pair = (event['mentee_id'], event['topic_id'])
        if pair in selected_pairs:
            groups[pair].append(event)
    results = []
    for pair, history in groups.items():
        successful = [e for e in history if e['date'] and (e['pc'] or e['mc'])]
        first_date = min((e['date'] for e in successful), default=None)
        if first_date:
            date_text = first_date.isoformat()
            if filters.get('date_from') and date_text < filters['date_from']:
                continue
            if filters.get('date_to') and date_text > filters['date_to']:
                continue
            # Require the first achievement to match activity filters too.
            # This prevents a later reassessment being called first competency.
            first = [e for e in successful if e['date'] == first_date]
            if not any(
                (not filters.get('mentor') or str(e['mentor_id']) == filters['mentor'])
                and (not filters.get('facility') or str(e['facility_id']) == filters['facility'])
                and (not filters.get('province') or str(e['province_id']) == filters['province'])
                and (not filters.get('thematic') or str(e['thematic_id']) == filters['thematic'])
                for e in first
            ):
                continue
        recorded_status = status_map.get(pair, {})
        stored_date = recorded_status.get('competent_date')
        earlier_status_date = bool(stored_date and first_date and stored_date < first_date)
        status_without_event = bool(recorded_status.get('status') == 'COMPETENT' and not first_date)
        missing_date = any(e['date'] is None for e in history)
        same_day = bool(first_date and any(e['ls'] and e['date'] == first_date for e in history))
        ls_visits = {
            e['visit_id'] for e in history
            if e['ls'] and e['date'] and (first_date is None or e['date'] < first_date)
        }
        if earlier_status_date:
            status = 'Stored competency predates PC/MC history — review required'
        elif status_without_event:
            status = 'Status says competent; dated PC/MC missing — review required'
        elif missing_date:
            status = 'Undated history — review required'
        elif same_day:
            status = 'Same-day sequence unknown — review required'
        elif first_date:
            status = 'First competency recorded'
        else:
            status = 'No dated competency recorded'
        results.append({
            'mentee_id': pair[0], 'topic_id': pair[1],
            'first_competency_date': first_date,
            'stored_competency_date': stored_date,
            'current_topic_status': {'NOT_STARTED': 'Not started', 'IN_PROGRESS': 'In progress', 'COMPETENT': 'Competent'}.get(recorded_status.get('status'), 'Not recorded'),
            'learning_sessions': len(ls_visits), 'status': status,
            'included': bool(first_date and not missing_date and not same_day and not earlier_status_date),
        })
    values = [r['learning_sessions'] for r in results if r['included']]
    return results, {
        'completed_pairs': len(values),
        'average_ls': round(mean(values), 2) if values else None,
        'median_ls': median(values) if values else None,
        'minimum_ls': min(values) if values else None,
        'maximum_ls': max(values) if values else None,
        'review_pairs': sum('review required' in r['status'] for r in results),
        'no_competency_pairs': sum(not r['first_competency_date'] for r in results),
    }, [{'sessions': n, 'pairs': count} for n, count in sorted(Counter(values).items())]


def build_dashboard_extensions(request, details_qs, scoped_history_qs, filters, province_id):
    """details_qs is already filtered; scoped_history_qs has authorization only."""
    if not request.user.is_superuser and province_id is None:
        details_qs = details_qs.none()
        scoped_history_qs = scoped_history_qs.none()

    # Current roster counts use province/facility. Activity selections additionally
    # restrict to unique mentees with a matching detail, avoiding join inflation.
    staff = Staff.objects.select_related('hfname__districtfk__provincefk')
    if not request.user.is_superuser:
        staff = staff.filter(hfname__districtfk__provincefk_id=province_id) if province_id is not None else staff.none()
    if filters.get('province'):
        staff = staff.filter(hfname__districtfk__provincefk_id=filters['province'])
    if filters.get('facility'):
        staff = staff.filter(hfname_id=filters['facility'])
    activity_selected = any(filters.get(k) for k in ('date_from', 'date_to', 'mentor', 'thematic'))
    if activity_selected:
        staff = staff.filter(pk__in=details_qs.order_by().values('menteename_id'))
    counts = staff.aggregate(
        active=Count('pk', filter=Q(status=True)),
        inactive=Count('pk', filter=Q(status=False)),
        unspecified=Count('pk', filter=Q(status__isnull=True)),
        total=Count('pk'),
    )
    facility_status = list(staff.order_by().values(
        facility_id=F('hfname_id'), facility=F('hfname__name'),
        province=F('hfname__districtfk__provincefk__name'),
    ).annotate(
        active=Count('pk', filter=Q(status=True)),
        inactive=Count('pk', filter=Q(status=False)),
        unspecified=Count('pk', filter=Q(status__isnull=True)), total=Count('pk'),
    ).order_by('province', 'facility', 'facility_id'))
    roster_rows = list(staff.order_by('hfname__districtfk__provincefk__name', 'hfname__name', 'hfname_id', 'firstname', 'pk').values(
        'id', 'firstname', 'lastname', 'fathername', 'status', 'gender',
        facility_id=F('hfname_id'), facility=F('hfname__name'),
        province=F('hfname__districtfk__provincefk__name'), position_name=F('position__name'),
    ))
    for row in roster_rows:
        row['mentee'] = ' '.join(part for part in (row['firstname'], row['lastname']) if part)
        row['status_label'] = 'Active' if row['status'] is True else 'Inactive' if row['status'] is False else 'Unspecified'
        row['gender_label'] = 'Female' if row['gender'] is True else 'Male' if row['gender'] is False else 'Unspecified'

    def aggregate(qs):
        return qs.annotate(
            mentees=Count('menteename_id', distinct=True),
            visits=Count('mentorshipvistfk_id', distinct=True),
            records=Count('pk'), ls=Count('pk', filter=Q(ls=True)),
            pc=Count('pk', filter=Q(pc=True)), mc=Count('pk', filter=Q(mc=True)),
        )

    gender_rows = list(aggregate(details_qs.order_by().values('menteename__gender').exclude(
        menteename_id__isnull=True)).order_by('menteename__gender'))
    visit_counts = mentor_visit_counts(details_qs)
    for row in gender_rows:
        value = row.pop('menteename__gender')
        row['mentor_visits'] = visit_counts.get('gender', {}).get(value, 0)
        row['gender'] = 'Female' if value is True else 'Male' if value is False else 'Unspecified'
    thematic_rows = list(aggregate(details_qs.order_by().exclude(thematicname_id__isnull=True).values(
        thematic_id=F('thematicname_id'), thematic=F('thematicname__name'),
    )).annotate(topics=Count('topicname_id', distinct=True)).order_by('thematic', 'thematic_id'))
    topic_rows = list(aggregate(details_qs.order_by().exclude(topicname_id__isnull=True).values(
        thematic_id=F('thematicname_id'), thematic=F('thematicname__name'),
        topic_id=F('topicname_id'), topic=F('topicname__name'),
    )).order_by('thematic', 'thematic_id', 'topic', 'topic_id'))
    for row in thematic_rows:
        row['mentor_visits'] = visit_counts.get('thematic', {}).get(row['thematic_id'], 0)
    for row in topic_rows:
        row['mentor_visits'] = visit_counts.get('topic', {}).get((row['thematic_id'], row['topic_id']), 0)
    definition_field = getattr(settings, 'MENTORSHIP_TOPIC_DEFINITION_FIELD', 'nameeng')
    MentorshipTopics._meta.get_field(definition_field)
    descriptions = {
        topic.pk: topic for topic in MentorshipTopics.objects.filter(
            pk__in=[r['topic_id'] for r in topic_rows])
    }
    for row in topic_rows:
        topic = descriptions[row['topic_id']]
        options = [(definition_field, getattr(topic, definition_field)), ('namedari', topic.namedari), ('namepashto', topic.namepashto)]
        field, text = next(((field, text) for field, text in options if text and str(text).strip()), ('', 'No topic description recorded'))
        row['definition'] = text
        row['definition_language'] = {'nameeng': 'English', 'namedari': 'Dari', 'namepashto': 'Pashto'}.get(field, 'Configured description') if field else ''
        row['thematic'] = row['thematic'] or 'Unassigned thematic area'

    pairs = set(details_qs.order_by().exclude(menteename_id__isnull=True).exclude(
        topicname_id__isnull=True).values_list('menteename_id', 'topicname_id'))
    # Entire authorized chronology, including sessions before the start filter.
    # Narrow by IDs in SQL and by exact selected pairs in the pure helper.
    history = scoped_history_qs.filter(
        menteename_id__in={p[0] for p in pairs}, topicname_id__in={p[1] for p in pairs},
    ).order_by().values(
        'ls', 'pc', 'mc', 'mentor_id',
        mentee_id=F('menteename_id'), topic_id=F('topicname_id'),
        date=F('mentorshipvistfk__visitdate'), visit_id=F('mentorshipvistfk_id'),
        province_id=F('mentorshipvistfk__facilityfk__districtfk__provincefk_id'),
        facility_id=F('mentorshipvistfk__facilityfk_id'),
        thematic_id=F('thematicname_id'), visit_round=F('mentorshipvistfk__visitround'),
    ).values('mentee_id', 'topic_id', 'date', 'visit_id', 'province_id',
             'facility_id', 'thematic_id', 'visit_round', 'ls', 'pc', 'mc', 'mentor_id')
    statuses = MenteeTopicStatus.objects.filter(
        mentee_id__in={p[0] for p in pairs}, topic_id__in={p[1] for p in pairs})
    if not request.user.is_superuser:
        statuses = statuses.filter(mentee__hfname__districtfk__provincefk_id=province_id) if province_id is not None else statuses.none()
    status_map = {(r['mentee_id'], r['topic_id']): r for r in statuses.values('mentee_id', 'topic_id', 'status', 'competent_date')}
    learning_rows, learning_stats, distribution = summarize_learning_history(history.iterator(chunk_size=2000), pairs, filters, status_map)
    names = {
        pk: ' '.join(part for part in (first, last) if part)
        for pk, first, last in Staff.objects.filter(pk__in={p[0] for p in pairs}).values_list('pk', 'firstname', 'lastname')
    }
    topics = dict(MentorshipTopics.objects.filter(pk__in={p[1] for p in pairs}).values_list('pk', 'name'))
    for row in learning_rows:
        row['mentee'] = names.get(row['mentee_id']) or str(row['mentee_id'])
        row['topic'] = topics.get(row['topic_id']) or str(row['topic_id'])
    quality = build_quality_data(request, details_qs, scoped_history_qs, filters, province_id, roster_rows, counts)
    for row in roster_rows:
        identity = quality['mentee_identity_map'][row['id']]
        row.update({key: identity[key] for key in ('mentee', 'fathername', 'fathername_missing')})
    for row in learning_rows:
        identity = quality['mentee_identity_map'][row['mentee_id']]
        row.update({key: identity[key] for key in ('mentee', 'fathername', 'fathername_missing')})
    for row in facility_status:
        known = row['active'] + row['inactive']
        row['attrition_rate'] = round(100 * row['inactive'] / known, 1) if known else None
    learning_rows.sort(key=lambda r: (r['mentee'], r['topic'], r['mentee_id'], r['topic_id']))
    return {
        **quality,
        'mentor_visit_counts': visit_counts,
        'mentee_status_kpis': counts, 'facility_status_rows': facility_status,
        'mentee_status_rows': roster_rows,
        'gender_rows': gender_rows, 'thematic_detail_rows': thematic_rows,
        'topic_detail_rows': topic_rows, 'learning_rows': learning_rows,
        'learning_stats': learning_stats, 'learning_distribution': distribution,
        'distinct_thematics_added': details_qs.exclude(thematicname_id__isnull=True).order_by().values('thematicname_id').distinct().count(),
        'distinct_topics_added': details_qs.exclude(topicname_id__isnull=True).order_by().values('topicname_id').distinct().count(),
        'mentee_status_scope_note': (
            'Current staff status; mentees with matching mentorship activity and a current facility matching the location filters.'
            if activity_selected else 'Current staff roster in the selected province/facility, including staff with no mentorship records.'
        ),
    }


def append_extension_sheets(workbook, data):
    """Call before the existing loop that styles every export sheet."""
    specifications = [
        ('Mentee_Status', 'facility_status_rows', [
            ('Province', 'province'), ('Facility ID', 'facility_id'), ('Facility', 'facility'),
            ('Active', 'active'), ('Inactive', 'inactive'), ('Unspecified', 'unspecified'), ('Total', 'total'), ('Attrition rate (current status) %', 'attrition_rate'),
        ]),
        ('Mentee_Roster', 'mentee_status_rows', [(key.title(), key) for key in ('province', 'facility_id', 'facility', 'id', 'mentee', 'fathername', 'position_name', 'gender_label', 'status_label')]),
        ('Gender', 'gender_rows', [(key.title(), key) for key in ('gender', 'mentees', 'mentor_visits', 'visits', 'records', 'ls', 'pc', 'mc')]),
        ('Thematic_Detail', 'thematic_detail_rows', [(key.title(), key) for key in ('thematic_id', 'thematic', 'mentees', 'topics', 'mentor_visits', 'visits', 'records', 'ls', 'pc', 'mc')]),
        ('Topic_Definitions', 'topic_detail_rows', [(key.title(), key) for key in ('thematic_id', 'thematic', 'topic_id', 'topic', 'definition', 'definition_language', 'mentees', 'mentor_visits', 'visits', 'records', 'ls', 'pc', 'mc')]),
        ('Learning_To_Competency', 'learning_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'topic_id', 'topic', 'learning_sessions', 'first_competency_date', 'stored_competency_date', 'current_topic_status', 'status', 'included')]),
        ('Name_Quality_Review', 'quality_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'facility', 'issues', 'recorded_values', 'suggested_values')]),
        ('Possible_Duplicates', 'duplicate_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'facility', 'other_id', 'other_mentee', 'other_fathername', 'other_facility', 'similarity', 'reason')]),
        ('Multiple_Facilities', 'multiple_facility_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'facility_count', 'facilities', 'first_seen', 'last_seen')]),
        ('Later_Facility_Records', 'facility_change_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'from_facility', 'previous_date', 'to_facility', 'later_date', 'reason')]),
        ('Facility_Sequence_Review', 'facility_sequence_review_rows', [(key.title(), key) for key in ('mentee_id', 'mentee', 'fathername', 'date', 'facilities', 'reason')]),
        ('Learning_Distribution', 'learning_distribution', [('LS visits before first competency', 'sessions'), ('Mentee-topic pairs', 'pairs')]),
    ]
    for title, data_key, columns in specifications:
        sheet = workbook.create_sheet(title)
        columns = [('MENTORSHIP VISITS (visit by mentors)' if key == 'mentor_visits' else 'Distinct visit IDs (reference)' if key == 'visits' else 'Father name' if key == 'fathername' else 'Father name B' if key == 'other_fathername' else label.replace('_', ' '), key) for label, key in columns]
        sheet.append([label for label, key in columns])
        if data_key == 'topic_detail_rows':
            sheet.sheet_properties.outlinePr.summaryBelow = False
            sheet.sheet_view.showOutlineSymbols = True
            summary_rows = []
            for area_id, grouped in groupby(data[data_key], key=lambda row: row['thematic_id']):
                rows = list(grouped)
                summary_row = [None] * len(columns)
                summary_row[0] = 'Thematic area'
                summary_row[1] = rows[0]['thematic']
                summary_row[3] = f'{len(rows)} topics — expand with +'
                sheet.append(summary_row)
                summary_index = sheet.max_row
                summary_rows.append(summary_index)
                sheet.row_dimensions[summary_index].collapsed = True
                first_detail = summary_index + 1
                for row in rows:
                    sheet.append([row.get(key) for label, key in columns])
                sheet.row_dimensions.group(first_detail, sheet.max_row, outline_level=1, hidden=True)
            sheet._msh_summary_rows = summary_rows
        else:
            for row in data[data_key]:
                sheet.append([row.get(key) for label, key in columns])
        # Names/definitions are data, even when they begin with '='.
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    notes = workbook.create_sheet('Extension_Methodology')
    notes.append(['Measure', 'Explanation'])
    notes.append(['Current staff status', data['mentee_status_scope_note']])
    notes.append(['Competency unit', 'Mentee-topic pair, first dated Patient Competent or Model Competent.'])
    notes.append(['Stored competency date', 'MenteeTopicStatus.competent_date is displayed for comparison. It has no visit/mentor reference; it is not used to invent an achievement event. An earlier stored date or competent status without a dated PC/MC event triggers review and exclusion from summaries.'])
    notes.append(['MENTORSHIP VISITS / visit by mentors', 'Unique visit date + mentor name (trimmed and lowercased), matching the original visit by mentors export. Missing names/dates excluded. Group counts can overlap across facilities, topics and provinces; do not sum them to obtain the global total.'])
    notes.append(['LS count', 'Distinct visits for the same mentee/topic strictly before first competency; repeated detail rows in one visit count once.'])
    notes.append(['History and filters', 'All authorized history contributes to LS counts. Filters select first achievement events, not subsequent reassessments. Earlier history outside the user province cannot be used.'])
    notes.append(['Excluded from summaries', 'Pairs with undated history or same-day LS/competency. Zero LS means no earlier LS recorded.'])
    notes.append(['Attrition (current-status proxy)', 'Inactive / (Active + Inactive) × 100. Unspecified status excluded; no exit dates or baseline cohort exist, so this is not a period dropout rate.'])
    notes.append(['Facility history', 'Same Staff ID observed in multiple authorized facilities through the selected end date. Later-facility events must match the filtered details; repeated rows are deduplicated by mentee/date/facility. Same-day multi-facility order is not inferred.'])
    notes.append(['Data quality', 'Review scope is selected current roster plus participants in matching visits. Name-case display is standardized without changing database values. Duplicate candidates are not confirmed duplicates.'])
    for key, value in data['learning_stats'].items():
        notes.append([key, value])


def format_dashboard_workbook(workbook, data):
    """MSH-inspired export styling, overview and native Excel topic outlines."""
    from datetime import date, datetime
    from math import ceil
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    green, orange, pale = '083D33', 'E36F1E', 'F2F7F4'
    overview = workbook.create_sheet('Dashboard_Overview', 0)
    overview.append(['Management Sciences for Health | Mentorship Dashboard'])
    overview.merge_cells('A1:C1')
    overview.append(['MNHIMS · IQoC-MNH | Filtered dashboard summary'])
    overview.merge_cells('A2:C2')
    overview.append([])
    overview.append(['Indicator', 'Value', 'Interpretation'])
    kpis = data['kpis']
    for label, value, note in [
        ('MENTORSHIP VISITS', kpis['mentorship_visits'], 'Visit by mentors: unique visit date + normalized mentor name; repeated details and visit IDs count once.'),
        ('Distinct visit IDs (reference)', kpis['distinct_visits'], 'Export-only reference count of distinct visit IDs; not the mentorship visit indicator.'),
        ('Distinct mentees mentored', kpis['distinct_mentees'], 'Distinct participants in matching mentorship details.'),
        ('Active mentees', data['mentee_status_kpis']['active'], data['mentee_status_scope_note']),
        ('Inactive mentees', data['mentee_status_kpis']['inactive'], data['mentee_status_scope_note']),
        ('Attrition rate (current status) %', data['attrition_rate'], 'Inactive / (Active + Inactive) × 100; current-status proxy, not period dropout.'),
        ('Mentees later recorded at another facility', data['quality_kpis']['later_facility_mentees'], 'Same Staff ID, strictly later dates; later event matches filters. Confirm reason with mentors.'),
        ('Mentees recorded in multiple facilities', data['quality_kpis']['multiple_facility_mentees'], 'Selected participants, full authorized history through selected end date.'),
        ('Father names not entered', data['quality_kpis']['missing_fathername'], 'Clinical mentor confirmation required.'),
        ('Possible duplicate pairs', data['quality_kpis']['duplicate_pairs'], 'Review candidates only; do not merge on name similarity alone.'),
        ('Unspecified status', data['mentee_status_kpis']['unspecified'], 'Current roster status is blank.'),
        ('Distinct thematic areas', data['distinct_thematics_added'], 'Distinct areas with matching mentorship details.'),
        ('Distinct topics', data['distinct_topics_added'], 'Distinct topics with matching mentorship details.'),
        ('Learning session details', kpis['total_ls'], 'LS detail instances; not unique mentees.'),
        ('Patient competent details', kpis['total_pc'], 'PC detail instances; not unique mentees.'),
        ('Model competent details', kpis['total_mc'], 'MC detail instances; not unique mentees.'),
        ('Average LS before competency', data['learning_stats']['average_ls'], 'Distinct earlier LS visits per completed mentee–topic pair.'),
        ('Median LS before competency', data['learning_stats']['median_ls'], 'Only pairs with usable chronology.'),
        ('Completed mentee–topic pairs', data['learning_stats']['completed_pairs'], 'Usable first dated PC/MC events selected by filters.'),
        ('Pairs requiring review', data['learning_stats']['review_pairs'], 'Incomplete or uncertain chronology excluded from statistics.'),
    ]:
        overview.append([label, value, note])
    for key, label in [('date_from', 'Date from'), ('date_to', 'Date to'), ('province', 'Province ID'), ('facility', 'Facility ID'), ('mentor', 'Mentor ID'), ('thematic', 'Thematic area ID')]:
        overview.append([label, data['filters'].get(key) or 'All', 'Applied dashboard filter'])
    overview.append(['Topic groups', 'Expand with +', 'Topic_Definitions starts with thematic summaries. Use Excel outline controls to reveal the underlying topic rows.'])
    overview.append(['Mentee topic lists', 'Expand columns', 'Mentee_Profile topic-list columns R:V are grouped and initially hidden. Use the outline controls to show them.'])

    edge = Side(style='thin', color='D5E3DC')
    for sheet in workbook.worksheets:
        header_row = 4 if sheet.title == 'Dashboard_Overview' else 1
        sheet.sheet_properties.tabColor = orange if sheet.title in ('Dashboard_Overview', 'Topic_Definitions') else green
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = f'A{header_row + 1}'
        sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        sheet.page_setup.orientation = 'landscape' if sheet.max_column > 5 else 'portrait'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = f'1:{header_row}'
        sheet.oddFooter.center.text = 'MSH | MNHIMS | Page &P of &N'
        sheet.oddFooter.center.size = 9
        for cell in sheet[header_row]:
            if cell.value in ('Visits', 'Visit Count'):
                cell.value = 'MENTORSHIP VISITS'
        widths = {}
        for column in sheet.iter_cols(min_row=header_row):
            header = str(column[0].value or '').lower()
            max_length = max((len(str(c.value)) for c in column if c.value is not None), default=10)
            if any(term in header for term in ('definition', 'explanation', 'interpretation', 'topics', 'status')):
                width = min(max(max_length + 2, 22), 48)
            else:
                width = min(max(max_length + 2, 13), 30)
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = width
            widths[column[0].column] = width
        summary_rows = set(getattr(sheet, '_msh_summary_rows', []))
        for cells in sheet.iter_rows(min_row=header_row):
            row_number = cells[0].row
            line_count = 1
            for cell in cells:
                cell.border = Border(bottom=edge)
                cell.font = Font(name='Calibri', size=11, color='203A32')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_number == header_row:
                    cell.fill = PatternFill('solid', fgColor=green)
                    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                elif row_number in summary_rows:
                    cell.fill = PatternFill('solid', fgColor='E4F0E8')
                    cell.font = Font(name='Calibri', size=11, bold=True, color=green)
                else:
                    cell.fill = PatternFill('solid', fgColor=pale if row_number % 2 == 0 else 'FFFFFF')
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                        if cell.value == 'FATHERNAME NOT ENTERED':
                            cell.font = Font(name='Calibri', size=11, bold=True, color='C62828')
                            cell.fill = PatternFill('solid', fgColor='FDE9E7')
                        if cell.value in ('Active', 'Inactive'):
                            cell.font = Font(name='Calibri', size=11, bold=True, color=green if cell.value == 'Active' else 'A44710')
                            cell.fill = PatternFill('solid', fgColor='DFF0E5' if cell.value == 'Active' else 'FFEBD8')
                    elif isinstance(cell.value, (date, datetime)):
                        cell.number_format = 'yyyy-mm-dd'
                    elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = '#,##0.##' if isinstance(cell.value, float) else '#,##0'
                        cell.alignment = Alignment(horizontal='right', vertical='top')
                if cell.value is not None:
                    line_count = max(line_count, sum(max(1, ceil(len(line) / widths[cell.column])) for line in str(cell.value).splitlines()))
            # Let wrapped descriptions occupy enough space when expanded.
            sheet.row_dimensions[row_number].height = max(30 if row_number == header_row else 24, 15 * line_count + 8)
        if sheet.title == 'Mentee_Profile':
            sheet.column_dimensions.group('R', 'V', outline_level=1, hidden=True)
            sheet.column_dimensions['W'].collapsed = True
            sheet.sheet_properties.outlinePr.summaryRight = True
        if sheet.title == 'Dashboard_Overview':
            sheet.column_dimensions['A'].width = 36
            sheet.column_dimensions['B'].width = 22
            sheet.column_dimensions['C'].width = 70
            sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color=green)
            sheet['A2'].font = Font(name='Calibri', size=11, color='A44710')
            sheet.row_dimensions[1].height = 30
            sheet.row_dimensions[2].height = 24
