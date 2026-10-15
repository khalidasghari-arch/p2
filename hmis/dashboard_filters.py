"""HMIS dashboard extension. No model/schema changes; inherits existing helpers."""
import os
import tempfile
from calendar import month_name
from decimal import Decimal, ROUND_HALF_UP
from django.contrib import messages
from django.db.models import Q, Sum, Count, Max, F
from django.core.paginator import Paginator
from django.http import FileResponse, HttpResponseRedirect
from django.template.response import TemplateResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from hmis.models import HMISFact, HMISMonthlySummary, IndicatorMetadata

# Must match slugify(indicator_name) in the supplied import pipeline.
SUMMARY_CODES = {
    'anc1': 'anc1', 'anc2': 'anc2', 'anc3': 'anc3', 'anc4': 'anc4',
    'pnc1': 'pnc1', 'pnc2': 'pnc2', 'n_delivery': 'n-delivery',
    'a_delivery': 'a-delivery', 'c_section': 'c-section',
    'lbw': 'lbw', 'stillbirth': 'stillbirth',
}
COMPOSITE_INDICATORS = (
    ('lbw', 'LBW (% of live births)'),
    ('babies-breastfed-1st-hour', 'Breastfed 1st Hour (% of live births)'),
    ('newborn-resuscitated', 'NB Resuscitated (% of live births)'),
    ('premature', 'Premature (% of live births)'),
)
LIVE_BIRTH_CODE = 'newbornalive'
IMPACT_INDICATORS = (
    ('neonatal', 'Facility neonatal mortality rate', 'total-number-of-neonatal-deaths',
     ('newbornalive',), 1000, 'per 1,000 live births'),
    ('maternal', 'Facility maternal mortality ratio', 'maternal-deaths-at-the-clinics',
     ('newbornalive',), 100000, 'per 100,000 live births'),
    ('stillbirth', 'Stillbirth rate', 'stillbirth',
     ('newbornalive', 'stillbirth'), 1000, 'per 1,000 total births'),
)
IMPACT_NOTE = (
    'For observed reporting facility-months, blank death, stillbirth and live-birth '
    'counts are treated as confirmed zeros under the project reporting rule. '
    'All valid positive birth denominators are retained, including months with '
    'blank death counts. A zero birth denominator cannot produce a rate. '
    'Totals are summed before calculating rates. Neonatal deaths use only '
    'Total-number-of-Neonatal-deaths; cause-specific death counts are not added '
    'to that total. Maternal deaths use Maternal-deaths-at-the-clinics. '
    'Live births use NewbornAlive. Total births = NewbornAlive + Stillbirth. '
    'These are facility-based monitoring measures, not population mortality '
    'estimates or proof of project impact. Verify the death definitions and '
    'coverage in the source reports, including referral/readmission cases and '
    'neonatal deaths within the first 28 days. Reduction requires a comparable '
    'baseline and follow-up period with consistent reporting.'
)


def calculate_impact(pairs):
    labels = {'newbornalive': 'Newborn Alive', 'stillbirth': 'Stillbirths',
              'total-number-of-neonatal-deaths': 'Total neonatal deaths',
              'maternal-deaths-at-the-clinics': 'Maternal deaths'}
    rows = []
    for key, label, numerator_code, denominator_codes, scale, unit in IMPACT_INDICATORS:
        numerator = denominator = Decimal('0')
        matched = zero_filled = 0
        issues = []
        for (facility, period), raw_values in sorted(pairs.items()):
            required = list(dict.fromkeys((numerator_code,) + denominator_codes))
            blank_zero = any(raw_values.get(code) is None for code in required)
            # The import pipeline omits blank cells; the user confirms these are zeros.
            values = {code: (raw_values.get(code) if raw_values.get(code) is not None else Decimal('0'))
                      for code in required}
            n = values[numerator_code]
            d = sum((values[code] for code in denominator_codes), Decimal('0'))
            included = False
            reason = action = ''
            if any(values[code] < 0 for code in required):
                reason = 'A required birth or death count is negative'
                action = 'Verify birth and death counts against the source register. Correct confirmed entry or import errors.'
            elif d == 0:
                reason = 'The birth denominator is zero; no rate can be calculated'
                action = ('Confirm there were no births; if the counts are correct, leave the rate unavailable.' if n == 0 else
                          'Verify birth reporting and whether deaths include referred or readmitted patients. Correct confirmed errors; do not invent a birth denominator.')
            else:
                matched += 1
                zero_filled += int(blank_zero)
                numerator += n
                denominator += d
                included = True
                if n > d:
                    reason = 'The death count exceeds the matched birth denominator'
                    action = 'Review both counts and their population coverage, including referral/readmission cases. Correct confirmed duplicates or reporting errors.'
            if reason:
                period = str(period)
                period_label = (f'{month_name[int(period[4:])]} {period[:4]}' if len(period)==6 and period.isdigit() and 1<=int(period[4:])<=12 else period)
                issues.append(dict(facility=facility, period=period_label, numerator=n, denominator=d,
                    numerator_display=f'{n:,.0f}' if n is not None else 'Not recorded',
                    denominator_display=f'{d:,.0f}' if d is not None else 'Not recorded',
                    reason=reason, action=action,
                    calculation_status='Included — review population coverage' if included else 'Not used in this rate'))
        rate = (numerator / denominator * scale).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) if matched else None
        excluded = sum(issue['calculation_status']=='Not used in this rate' for issue in issues)
        row = dict(code=key, label=label, unit=unit, scale=scale,
            numerator=numerator if matched else None, denominator=denominator if matched else None,
            numerator_display=f'{numerator:,.0f}' if matched else '—',
            denominator_display=f'{denominator:,.0f}' if matched else '—',
            rate=rate, rate_display=f'{rate:,.2f}' if rate is not None else '—',
            matched=matched, excluded=excluded, zero_filled=zero_filled, issues=issues, issue_count=len(issues),
            denominator_label='Total births (live + stillbirths)' if key=='stillbirth' else 'Live births',
            formula=f"{labels[numerator_code]} ÷ {'(Newborn Alive + Stillbirths)' if key=='stillbirth' else 'Newborn Alive'} × {scale:,}",
            status=(f'{len(issues)} facility-month checks — view details' if issues else
                    (f'Confirmed blank zeros included for {zero_filled} facility-month' + ('s' if zero_filled != 1 else '') if zero_filled else 'No issues detected in matched records') if matched else 'No rate available — check imported data'))
        rows.append(row)
    return rows


COMPOSITE_NOTE = (
    'Each percentage = the recorded numerator ÷ Newborn Alive × 100 for matching '
    'facility-months. Only pairs with a nonblank, nonnegative numerator and a '
    'positive denominator contribute to both totals. Recorded zero numerators '
    'are included; missing values are not treated as zero. Combined rates use '
    'summed counts, not an average of percentages. These measures use the selected '
    'location and dates and remain available regardless of the indicator dropdown. '
    'NB Resuscitated is a share of live births, not a resuscitation success rate '
    'or coverage among babies needing resuscitation. Numerators must refer to '
    'the same live-born population as Newborn Alive; out-of-facility or referred '
    'cases can make that comparison inappropriate.'
)


def calculate_composites(pairs):
    """Rates plus facility/month follow-up, without changing eligibility rules."""
    result = []
    for code, label in COMPOSITE_INDICATORS:
        short_name = label.split(' (')[0]
        numerator = denominator = Decimal('0')
        matched = excluded = above = 0
        missing_numerator = missing_denominator = zero_denominator = negative = 0
        issues = []
        for (facility, periodcode), values in sorted(pairs.items()):
            n, d = values.get(code), values.get(LIVE_BIRTH_CODE)
            reason = action = ''
            included = False
            if n is None or d is None:
                excluded += 1
                missing_numerator += n is None
                missing_denominator += d is None
                if n is None and d is None:
                    reason = f'{short_name} and Newborn Alive values are missing from imported data'
                elif n is None:
                    reason = f'{short_name} value is missing from imported data'
                else:
                    reason = 'Newborn Alive value is missing from imported data'
                action = 'Check the facility report for this month. Confirm the missing value; record zero only if verified, then import the corrected report.'
            elif n < 0 or d < 0:
                excluded += 1
                negative += 1
                reason = 'A recorded count is negative'
                action = 'Compare the counts with the source register. Correct any entry or import error and re-import the report.'
            elif d == 0:
                excluded += 1
                zero_denominator += 1
                if n == 0:
                    reason = 'Newborn Alive is zero; a percentage cannot be calculated'
                    action = 'Confirm that no live births occurred. If both counts are correct, leave the percentage unavailable; no correction is needed.'
                else:
                    reason = f'{short_name} is recorded, but Newborn Alive is zero'
                    action = 'Check the live-birth count and whether the indicator includes referred or outborn babies. Correct confirmed errors and re-import.'
            else:
                numerator += n
                denominator += d
                matched += 1
                included = True
                if n > d:
                    above += 1
                    reason = f'{short_name} count exceeds Newborn Alive'
                    action = 'Compare both counts with the source register. Check duplicate entries and whether both counts cover the same babies. Correct confirmed errors and re-import.'
            if reason:
                period = str(periodcode)
                period_label = (f'{month_name[int(period[4:])]} {period[:4]}'
                                if len(period) == 6 and period.isdigit() and 1 <= int(period[4:]) <= 12 else period)
                issues.append(dict(facility=facility, periodcode=period, period=period_label,
                    numerator=n, denominator=d,
                    numerator_display=f'{n:,.0f}' if n is not None else 'Not recorded',
                    denominator_display=f'{d:,.0f}' if d is not None else 'Not recorded',
                    reason=reason, action=action, included=included,
                    calculation_status='Included — count review required' if included else 'Not used in this percentage'))
        rate = (numerator / denominator * 100).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) if matched else None
        messages = []
        if not matched:
            messages.append('No percentage available')
        if excluded:
            unit = 'facility-month' if excluded == 1 else 'facility-months'
            messages.append(f'{excluded} {unit} not used — check missing, zero or invalid counts')
        if above:
            unit = 'facility-month' if above == 1 else 'facility-months'
            messages.append(f'{above} {unit} needs count review')
        if not pairs:
            messages.append('Check whether the newborn indicators were imported for this selection')
        result.append(dict(code=code, label=label,
            numerator=numerator if matched else None, denominator=denominator if matched else None,
            rate=rate, rate_display=f'{rate:.2f}%' if rate is not None else '—',
            numerator_display=f'{numerator:,.0f}' if matched else '—',
            denominator_display=f'{denominator:,.0f}' if matched else '—',
            matched=matched, excluded=excluded, above=above,
            missing_numerator=missing_numerator, missing_denominator=missing_denominator,
            zero_denominator=zero_denominator, negative=negative,
            issues=issues, issue_count=len(issues),
            status='; '.join(messages) or 'No issues detected in matched records',
            warning=bool(issues)))
    return result


BASELINE_NOTE = (
    'Baseline: distinct facilities observed anywhere in HMIS history, limited by '
    'the selected facility group and location. This assumes those facilities are '
    'expected in every displayed month; facilities never recorded and changes '
    'in operational status cannot be determined from these models. A monthly '
    'summary with at least one nonblank indicator, or a nonblank HMIS fact, counts '
    'as reported. Zero counts as reported. Missing means no recorded report, '
    'not a confirmed overdue submission.'
)


def month_periods(start, end, years=(), months=()):
    """Include gaps inside the database's observed calendar range."""
    if not start or not end:
        return []
    year, month = int(start[:4]), int(start[4:])
    result = []
    while f'{year:04d}{month:02d}' <= end:
        if (not years or year in years) and (not months or month in months):
            result.append(f'{year:04d}{month:02d}')
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return result


def reporting_rows(roster, submissions, periods):
    """One reporting count per facility/month, independent of indicator selection."""
    submitted = {}
    for hf, period in submissions:
        submitted.setdefault(period, set()).add(hf)
    expected = set(roster)
    rows = []
    for period in periods:
        reporting = expected & submitted.get(period, set())
        missing = sorted(expected - reporting, key=str.casefold)
        rows.append({
            'periodcode': period,
            'period': f'{month_name[int(period[4:])]} {period[:4]}',
            'expected': len(expected), 'reported': len(reporting),
            'missing_count': len(missing),
            'rate': round(100 * len(reporting) / len(expected), 2) if expected else None,
            'rate_display': f'{100 * len(reporting) / len(expected):.1f}%' if expected else '—',
            'missing': [dict(roster[hf], hf=hf) for hf in missing],
        })
    return rows


class HMISDashboardFilterMixin:
    """Place before admin.ModelAdmin in HMISDashboardAdmin's inheritance list."""

    def _catalogue(self):
        fields = ('prov', 'dist', 'hf', 'hiva_hfs', 'year', 'month', 'periodcode')
        # Union includes facilities/periods present in either reporting table.
        rows = list(HMISMonthlySummary.objects.order_by().values_list(*fields).union(
            HMISFact.objects.order_by().values_list(*fields)))
        result = []
        for prov, dist, hf, hiva, year, month, period in rows:
            period = str(period or '')
            if len(period) != 6 or not period.isdigit() or not 1 <= int(period[4:]) <= 12:
                continue
            if year != int(period[:4]) or month != int(period[4:]):
                continue
            result.append(dict(province=prov, district=dist, facility=hf,
                               group='hiva' if hiva else 'non_hiva',
                               year=year, month=month, periodcode=period))
        return result

    @staticmethod
    def _geo_rows(rows, filters):
        return [r for r in rows if
                (filters['facility_group'] == 'all' or r['group'] == filters['facility_group'])
                and (not filters['prov'] or r['province'] == filters['prov'])
                and (not filters['dist'] or r['district'] == filters['dist'])
                and (not filters['hf'] or r['facility'] == filters['hf'])]

    def _read_filters(self, request):
        def integers(key, lower, upper):
            values = []
            for raw in request.GET.getlist(key):
                try:
                    value = int(raw)
                except (TypeError, ValueError):
                    continue
                if lower <= value <= upper and value not in values:
                    values.append(value)
            return sorted(values)
        # Use actual fact codes for every indicator; preserve old summary URLs.
        indicators = list(dict.fromkeys(
            SUMMARY_CODES.get(value.strip(), value.strip())
            for value in request.GET.getlist('indicator')
            if value.strip() and len(value.strip()) <= 128))
        group = request.GET.get('facility_group', 'all')
        return dict(prov=request.GET.get('prov', '').strip(),
                    dist=request.GET.get('dist', '').strip(), hf=request.GET.get('hf', '').strip(),
                    facility_group=group if group in ('all', 'hiva', 'non_hiva') else 'all',
                    year=integers('year', 1, 9999), month=integers('month', 1, 12),
                    indicator=indicators)

    def _filter_queryset(self, queryset, filters):
        for name in ('prov', 'dist', 'hf'):
            if filters[name]:
                queryset = queryset.filter(**{name: filters[name]})
        for name in ('year', 'month'):
            if filters[name]:
                queryset = queryset.filter(**{name + '__in': filters[name]})
        return self._apply_facility_group(queryset, filters['facility_group'])

    def _filter_fact_queryset(self, queryset, filters):
        queryset = self._filter_queryset(queryset, filters)
        if filters['indicator']:
            queryset = queryset.filter(indicator_code__in=filters['indicator'])
        return queryset

    def _filter_options(self, rows, filters):
        groups = sorted({r['group'] for r in rows})
        scope = [r for r in rows if filters['facility_group'] == 'all' or r['group'] == filters['facility_group']]
        provinces = sorted({r['province'] for r in scope if r['province']})
        if filters['prov']:
            scope = [r for r in scope if r['province'] == filters['prov']]
        districts = sorted({r['district'] for r in scope if r['district']})
        if filters['dist']:
            scope = [r for r in scope if r['district'] == filters['dist']]
        facilities = sorted({r['facility'] for r in scope if r['facility']})
        if filters['hf']:
            scope = [r for r in scope if r['facility'] == filters['hf']]
        years = sorted({r['year'] for r in scope}, reverse=True)
        if filters['year']:
            scope = [r for r in scope if r['year'] in filters['year']]
        months = sorted({r['month'] for r in scope})
        return dict(groups=[{'value': x, 'label': 'HIVA Health Facilities' if x == 'hiva'
                            else 'Non-HIVA Health Facilities'} for x in groups],
                    provinces=provinces, districts=districts, facilities=facilities,
                    years=years, months=[{'value': x, 'label': month_name[x]} for x in months],
                    location_cascade=rows)

    def _indicator_catalogue(self, queryset):
        """All recorded fact indicators, including zero values and future additions."""
        records = list(queryset.order_by().exclude(indicator_code='')
                       .filter(value__isnull=False).values('indicator_code')
                       .annotate(recorded_name=Max('indicator_name')))
        codes = [record['indicator_code'] for record in records]
        metadata = {}
        for item in IndicatorMetadata.objects.filter(indicator_code__in=codes).order_by(
                '-is_active', 'sort_order', 'indicator_name'):
            metadata.setdefault(item.indicator_code, item)
        options = []
        for record in records:
            code = record['indicator_code']
            item = metadata.get(code)
            label = (item.indicator_name if item and item.indicator_name else
                     record['recorded_name'] or code)
            order = item.sort_order if item and item.sort_order is not None else 1_000_000
            options.append((order, label.casefold(), code, label))
        options.sort(key=lambda item: (item[0], item[1], item[2]))
        return [(code, label) for _order, _sort_label, code, label in options]

    def _indicator_options(self, filters):
        # Apply location/date scope, but never restrict by the selected indicators.
        return self._indicator_catalogue(
            self._filter_queryset(HMISFact.objects.all(), filters))

    def _selection_labels(self, filters):
        labels = dict(self._indicator_catalogue(HMISFact.objects.filter(
            indicator_code__in=filters['indicator'])))
        return [labels.get(code, code) for code in filters['indicator']]

    def _reporting(self, catalogue, filters):
        # Latest observed location/group per facility; dates never shrink the roster.
        latest = {}
        for r in sorted(catalogue, key=lambda r: (r['periodcode'], r['province'], r['district'], r['group'])):
            if r['facility']:
                latest[r['facility']] = r
        roster = {r['facility']: {'prov': r['province'], 'dist': r['district'],
                                 'group': r['group']} for r in self._geo_rows(list(latest.values()), filters)}
        observed = sorted({r['periodcode'] for r in catalogue})
        periods = month_periods(observed[0], observed[-1], filters['year'], filters['month']) if observed else []
        nonblank = Q()
        for field in SUMMARY_CODES:
            nonblank |= Q(**{field + '__isnull': False})
        # Query by roster identity, regardless of historical group/location changes.
        summaries = HMISMonthlySummary.objects.filter(nonblank, hf__in=roster, periodcode__in=periods)
        facts = HMISFact.objects.filter(hf__in=roster, periodcode__in=periods, value__isnull=False)
        submissions = summaries.order_by().values_list('hf', 'periodcode').union(
            facts.order_by().values_list('hf', 'periodcode'))
        return reporting_rows(roster, submissions, periods)

    def _composite_data(self, filters):
        codes = [code for code, _label in COMPOSITE_INDICATORS] + [LIVE_BIRTH_CODE]
        facts = self._filter_queryset(HMISFact.objects.all(), filters).filter(indicator_code__in=codes)
        grouped = facts.order_by().values('hf', 'periodcode', 'indicator_code').annotate(
            total=Sum('value'), nonblank=Count('value'))
        pairs = {}
        monthly_pairs = {}
        for record in grouped.iterator(chunk_size=2000):
            key = (record['hf'], record['periodcode'])
            value = record['total'] if record['nonblank'] else None
            pairs.setdefault(key, {})[record['indicator_code']] = value
            monthly_pairs.setdefault(record['periodcode'], {}).setdefault(key, {})[record['indicator_code']] = value
        monthly = []
        for period in sorted(monthly_pairs):
            for row in calculate_composites(monthly_pairs[period]):
                monthly.append(dict(row, periodcode=period, period=self._period_label(period)))
        return dict(rows=calculate_composites(pairs), monthly=monthly, note=COMPOSITE_NOTE)

    def _impact_data(self, filters):
        codes = {'total-number-of-neonatal-deaths', 'maternal-deaths-at-the-clinics', 'newbornalive', 'stillbirth'}
        source_facts = self._filter_queryset(HMISFact.objects.all(), filters)
        summaries = self._filter_queryset(HMISMonthlySummary.objects.all(), filters)
        # Observe reports across all fields, not only death fields; never create reports
        # for roster-only facility-months with no stored reporting record.
        observed = summaries.order_by().values_list('hf','periodcode').union(
            source_facts.order_by().values_list('hf','periodcode'))
        pairs = {key: {} for key in observed}
        monthly_pairs = {}
        for key in pairs:
            monthly_pairs.setdefault(key[1], {})[key] = pairs[key]
        for record in source_facts.filter(indicator_code__in=codes).order_by().values(
                'hf','periodcode','indicator_code').annotate(
                total=Sum('value'), nonblank=Count('value')).iterator(chunk_size=2000):
            key = (record['hf'], record['periodcode'])
            value = record['total'] if record['nonblank'] else None
            pairs.setdefault(key,{})[record['indicator_code']] = value
            monthly_pairs.setdefault(record['periodcode'],{}).setdefault(key,pairs[key])[record['indicator_code']] = value
        monthly = [dict(row,periodcode=period,period=self._period_label(period))
                   for period in sorted(monthly_pairs)
                   for row in calculate_impact(monthly_pairs[period])]
        return dict(rows=calculate_impact(pairs),monthly=monthly,note=IMPACT_NOTE)

    @staticmethod
    def _msh_header(worksheet, headers, widths=None):
        worksheet.freeze_panes = 'A2'
        color = 'EF6C00' if worksheet.title == 'Missing Reports' else '073F36'
        cells = []
        for index, header in enumerate(headers, 1):
            cell = WriteOnlyCell(worksheet, value=header)
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(color='FFFFFF', bold=True, size=11)
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            cells.append(cell)
            worksheet.column_dimensions[get_column_letter(index)].width = widths[index - 1] if widths else 20
        worksheet.append(cells)

    def _multi_excel(self, facts, filters, reporting, composite, impact):
        workbook = Workbook(write_only=True)
        scope = workbook.create_sheet('Summary')
        title = WriteOnlyCell(scope, value='MSH | IQoC-MNH | HMIS filtered export')
        title.font = Font(color='FFFFFF', bold=True, size=14)
        title.fill = PatternFill('solid', fgColor='073F36')
        scope.append([title, timezone.localtime().isoformat()])
        for key in ('facility_group', 'prov', 'dist', 'hf', 'year', 'month', 'indicator'):
            values = filters[key]
            if key == 'month':
                values = [month_name[x] for x in values]
            if key == 'indicator':
                values = self._selection_labels(filters)
            scope.append([key, self._excel_text(', '.join(map(str, values)) if isinstance(values, list)
                                               else values or 'All') or 'All'])
        scope.append(['Reporting baseline', BASELINE_NOTE])
        self._msh_header(scope, ('Indicator Code', 'Indicator Name', 'Total', 'Nonblank Values'), (20, 48, 18, 18))
        for r in facts.values('indicator_code').annotate(name=Max('indicator_name'), total=Sum('value'), n=Count('value')).order_by('indicator_code').iterator():
            scope.append([self._excel_text(r['indicator_code']), self._excel_text(r['name']), r['total'], r['n']])
        details = workbook.create_sheet('Details')
        fields = ('id', 'source_upload_id', 'prov', 'dist', 'hf', 'periodcode', 'year', 'month',
                  'month_name', 'hf_name_cleaned', 'hfid', 'hiva_hfs', 'indicator_code', 'indicator_name', 'value', 'created_at')
        self._msh_header(details, fields, (12,16,18,18,34,14,10,10,14,34,12,12,20,48,18,24))
        for row in facts.order_by().values_list(*fields).iterator(chunk_size=5000):
            values = [self._excel_text(v) for v in row]
            values[-1] = self._datetime_display(values[-1])
            details.append(values)
        monthly = workbook.create_sheet('Reporting Completeness')
        self._msh_header(monthly, ('Month', 'Expected Known Facilities', 'Reporting', 'Missing', 'Completeness (%)'), (24,28,16,16,24))
        missing = workbook.create_sheet('Missing Reports')
        self._msh_header(missing, ('Month', 'Province', 'District', 'Facility', 'Facility Group'), (24,22,22,40,20))
        for row in reporting:
            monthly.append([row['period'], row['expected'], row['reported'], row['missing_count'], row['rate']])
            for facility in row['missing']:
                missing.append([self._excel_text(x) for x in (row['period'], facility['prov'], facility['dist'], facility['hf'], facility['group'])])
        composite_sheet = workbook.create_sheet('Composite Indicators')
        headers = ('Scope', 'Indicator', 'Numerator', 'Newborn Alive (matched)', 'Rate (%)',
                   'Matched Facility-Months', 'Excluded Facility-Months', 'Status')
        self._msh_header(composite_sheet, headers, (24,44,18,26,18,26,26,65))
        for row in composite['rows']:
            composite_sheet.append([self._excel_text(x) for x in (
                'Selected period', row['label'], row['numerator'], row['denominator'], row['rate'],
                row['matched'], row['excluded'], row['status'])])
        for row in composite['monthly']:
            composite_sheet.append([self._excel_text(x) for x in (
                row['period'], row['label'], row['numerator'], row['denominator'], row['rate'],
                row['matched'], row['excluded'], row['status'])])
        composite_sheet.append([])
        composite_sheet.append(['Calculation', COMPOSITE_NOTE])
        checks = workbook.create_sheet('Composite Data Checks')
        self._msh_header(checks,
            ('Indicator', 'Facility', 'Month', 'Indicator Count', 'Newborn Alive',
             'Issue', 'Suggested Action', 'Calculation Treatment'),
            (38,42,24,20,20,65,90,38))
        for row in composite['rows']:
            for issue in row['issues']:
                checks.append([self._excel_text(value) for value in (
                    row['label'], issue['facility'], issue['period'], issue['numerator'],
                    issue['denominator'], issue['reason'], issue['action'], issue['calculation_status'])])
        impact_sheet = workbook.create_sheet('Impact Indicators')
        self._msh_header(impact_sheet, ('Scope','Indicator','Death Count','Matched Birth Denominator','Rate','Unit','Matched Facility-Months','Data Check'),
                         (24,42,18,28,18,28,26,65))
        for row in impact['rows']:
            impact_sheet.append([self._excel_text(value) for value in (
                'Selected period',row['label'],row['numerator'],row['denominator'],row['rate'],row['unit'],row['matched'],row['status'])])
        for row in impact['monthly']:
            impact_sheet.append([self._excel_text(value) for value in (
                row['period'],row['label'],row['numerator'],row['denominator'],row['rate'],row['unit'],row['matched'],row['status'])])
        impact_sheet.append([])
        impact_sheet.append(['Methodology',IMPACT_NOTE])
        for row in impact['rows']:
            impact_sheet.append([row['label'],row['formula']])
        impact_checks = workbook.create_sheet('Impact Data Checks')
        self._msh_header(impact_checks, ('Indicator','Facility','Month','Death Count','Birth Denominator','Issue','Suggested Action','Calculation Treatment'),
                         (42,42,24,18,24,65,90,40))
        for row in impact['rows']:
            for issue in row['issues']:
                impact_checks.append([self._excel_text(value) for value in (
                    row['label'],issue['facility'],issue['period'],issue['numerator'],issue['denominator'],issue['reason'],issue['action'],issue['calculation_status'])])
        path = None
        try:
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as handle:
                path = handle.name
            workbook.save(path)
            response = FileResponse(open(path, 'rb'), as_attachment=True,
                                    filename='HMIS_Filtered_Dashboard.xlsx',
                                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Cache-Control'] = 'no-store'
            response._resource_closers.append(lambda: os.path.exists(path) and os.remove(path))
            return response
        except Exception:
            if path and os.path.exists(path):
                os.remove(path)
            raise

    @staticmethod
    def _page_link(request, key, number, anchor):
        params = request.GET.copy()
        params.pop('print', None)
        params.pop('export', None)
        params[key] = number
        if anchor == 'hmis-results':
            params['detail_tab'] = key.replace('_page','')+'-panel'
        return '?' + params.urlencode() + '#' + anchor

    def _page(self, request, objects, key, size, anchor, complete=False):
        paginator = Paginator(objects, size)
        page = paginator.get_page(request.GET.get(key, '1'))
        info = dict(number=page.number, pages=paginator.num_pages, total=paginator.count,
                    start=page.start_index(), end=page.end_index(),
                    previous=self._page_link(request,key,page.previous_page_number(),anchor) if page.has_previous() else '',
                    next=self._page_link(request,key,page.next_page_number(),anchor) if page.has_next() else '')
        if complete:
            info.update(number=1,pages=1,start=1 if paginator.count else 0,end=paginator.count,previous='',next='')
            return list(objects), info
        return list(page.object_list), info

    def _paged_detail_tables(self, request, queryset, cards, selected, complete=False):
        """Aggregate only the geography/month groups displayed on this page."""
        codes = [card['indicator_code'] for card in cards]
        indexes = {code: index for index,code in enumerate(codes)}
        output = {'detailed_indicator_count':len(cards)}
        search = request.GET.get('facility_search','').strip()[:200]
        for kind, fields, size in (
            ('monthly',('periodcode',),24),
            ('province',('prov',),25),
            ('facility',('prov','hf','hiva_hfs'),25)):
            scope = queryset
            if kind == 'facility' and search and not complete:
                scope = scope.filter(Q(hf__icontains=search)|Q(prov__icontains=search))
            groups = scope.order_by().values(*fields)
            if kind != 'monthly' and selected in indexes:
                groups = groups.annotate(sort_total=Sum('value',filter=Q(indicator_code=selected))).order_by(
                    F('sort_total').desc(nulls_last=True),*fields)
            else:
                groups = groups.distinct().order_by(*fields)
            chosen, pager = self._page(request,groups,kind+'_page',size,'hmis-results',complete)
            keys = [tuple(group[field] for field in fields) for group in chosen]
            rows = {key:dict(key=key,raw_values=[None]*len(codes)) for key in keys}
            condition = Q(pk__in=[])
            for key in keys:
                condition |= Q(**dict(zip(fields,key)))
            detail_scope = scope if complete else scope.filter(condition)
            grouped = detail_scope.order_by().values(*fields,'indicator_code').annotate(
                total=Sum('value'), nonblank=Count('value'))
            for record in grouped.iterator(chunk_size=1000):
                code = record['indicator_code'] or 'Not coded'
                index = indexes.get(code)
                if index is not None:
                    key = tuple(record[field] for field in fields)
                    rows[key]['raw_values'][index] = float(record['total']) if record['nonblank'] and record['total'] is not None else None
            for key,row in rows.items():
                row['indicator_values'] = [dict(raw=value,display=self._number_display(value)) for value in row['raw_values']]
                if kind == 'monthly':
                    row.update(periodcode=key[0],period=self._period_label(key[0]))
                elif kind == 'province':
                    row['province']=key[0] or 'Not specified'
                else:
                    row.update(province=key[0] or 'Not specified',facility=key[1] or 'Not specified',
                               hiva_hfs=bool(key[2]),facility_group='HIVA HF' if key[2] else 'Non-HIVA HF')
            output['detailed_'+kind+'_rows']=list(rows.values())
            output[kind+'_pager']=pager
        output['facility_search']=search
        output['facility_search_params']=[(key,value) for key,values in request.GET.lists()
            if key not in ('facility_search','facility_page','detail_tab','print','export') for value in values]
        requested_tab=request.GET.get('detail_tab')
        output['detail_active_tab']=(requested_tab if requested_tab in ('monthly-panel','province-panel','facility-panel') else
                                    'facility-panel' if search or request.GET.get('facility_page') else
                                    'province-panel' if request.GET.get('province_page') else 'monthly-panel')
        return output

    def _screen_details(self, request, reporting, composite, impact, complete=False):
        """Copies for presentation: exports and full totals keep the complete lists."""
        report_page, report_pager = self._page(request,reporting,'reporting_page',12,'hmis-reporting',complete)
        visible_reporting=[]
        for row in report_page:
            missing, pager = self._page(request,row['missing'],'missing_'+row['periodcode'],10,'hmis-reporting',complete)
            visible_reporting.append(dict(row,missing=missing,missing_pager=pager))
        visible_metrics={}
        for name,source in (('composite',composite),('impact',impact)):
            rows=[]
            for row in source['rows']:
                issues,pager=self._page(request,row['issues'],'checks_'+name+'_'+row['code'],10,
                                       ('composite-check-' if name=='composite' else 'impact-check-')+row['code'],complete)
                rows.append(dict(row,issues=issues,issue_pager=pager))
            monthly,pager=self._page(request,source['monthly'],name+'_month_page',24,
                                    'hmis-composites' if name=='composite' else 'hmis-impact',complete)
            visible_metrics[name]=dict(source,rows=rows,monthly=monthly,monthly_pager=pager)
        return dict(reporting_rows=visible_reporting,reporting_pager=report_pager,**visible_metrics)

    @staticmethod
    def _print_cost(facts, cards, reporting, composite, impact):
        group_rows=sum(facts.order_by().values(*fields).distinct().count() for fields in
                       (('periodcode',),('prov',),('prov','hf','hiva_hfs')))
        missing=sum(row['missing_count'] for row in reporting)
        checks=sum(len(row['issues']) for data in (composite,impact) for row in data['rows'])
        monthly=sum(len(data['monthly']) for data in (composite,impact))
        return group_rows*max(1,len(cards))+missing*4+(checks+monthly)*8

    def changelist_view(self, request, extra_context=None):
        if not self.has_view_permission(request):
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        filters = self._read_filters(request)
        catalogue = self._catalogue()
        options = self._filter_options(catalogue, filters)
        indicator_options = self._indicator_options(filters)
        summaries = self._filter_queryset(HMISMonthlySummary.objects.all(), filters)
        facts = self._filter_fact_queryset(HMISFact.objects.all(), filters)
        reporting = self._reporting(catalogue, filters)
        composite = self._composite_data(filters)
        impact = self._impact_data(filters)
        if request.GET.get('export') == 'xlsx':
            if facts.count() > self._filtered_excel_row_limit():
                self.message_user(request, 'Please narrow the filters: the Excel detail row limit was exceeded.', level=messages.WARNING)
                params = request.GET.copy()
                params.pop('export', None)
                return HttpResponseRedirect(request.path + '?' + params.urlencode())
            return self._multi_excel(facts, filters, reporting, composite, impact)
        first = filters['indicator'][0] if filters['indicator'] else 'anc1'
        summary_field = next((field for field, code in SUMMARY_CODES.items() if code == first), 'anc1')
        data = self._build_dashboard_data(summaries, summary_field)
        data['selected_indicator_label'] = ', '.join(self._selection_labels(filters)) or 'All HMIS indicators'
        cards, charts = self._all_hmis_indicator_trends(facts)
        data['all_hmis_indicator_cards'] = cards
        data['chart_data']['all_hmis_indicator_trends'] = charts
        for_print = bool(request.GET.get('print'))
        if for_print and self._print_cost(facts,cards,reporting,composite,impact) > 5000:
            self.message_user(request,
                'PDF was not started because this selection contains too many detail rows. '
                'Select fewer months, facilities or indicators, or use Excel for the full filtered data.',
                level=messages.WARNING)
            params=request.GET.copy()
            params.pop('print',None)
            return HttpResponseRedirect(request.path+'?'+params.urlencode())
        data.update(self._paged_detail_tables(request,facts,cards,first,complete=for_print))
        expected = sum(r['expected'] for r in reporting)
        reported = sum(r['reported'] for r in reporting)
        data['kpis'].insert(0, dict(label='Monthly Reporting Completeness',
            value=self._percent_display(self._percent(reported, expected)),
            help=f'{reported:,}/{expected:,} expected known facility-month reports', **{'class': 'green'}))
        for kpi in data['kpis']:
            if kpi['label'] == 'Data Completeness':
                kpi['label'] = 'Summary Field Completeness'
        for_print = bool(request.GET.get('print'))
        export_params = request.GET.copy()
        export_params.pop('print', None)
        export_params['export'] = 'xlsx'
        print_params = request.GET.copy()
        print_params.pop('export', None)
        print_params['print'] = 'filtered'
        primary_names = ('Monthly Reporting Completeness', 'Health Facilities', 'Reporting Periods',
                         'ANC 1st Visit', 'PNC 1st Visit', 'Total Deliveries')
        data['primary_kpis'] = [kpi for name in primary_names for kpi in data['kpis'] if kpi['label'] == name]
        data['additional_kpis'] = [kpi for kpi in data['kpis'] if kpi['label'] not in primary_names]
        context = {**self.admin_site.each_context(request), **data,
                   'title': 'HMIS Performance Dashboard', 'opts': self.model._meta,
                   'filters': filters, 'filter_options': options, 'indicator_options': indicator_options,
                   **self._screen_details(request,reporting,composite,impact,complete=for_print),
                   'reporting_note': BASELINE_NOTE,
                   'export_url': '?' + export_params.urlencode(), 'print_url': '?' + print_params.urlencode(),
                   'excel_export_limit_display': f'{self._filtered_excel_row_limit():,}',
                   'print_all_mode': for_print}
        if extra_context:
            context.update(extra_context)
        request.current_app = self.admin_site.name
        return TemplateResponse(request, self.change_list_template, context)
